"""Data-center economics figures from the post-processed tables.

Same data-center figure style (Times New Roman, capitalized region names) as
scripts/plot_dc_tables.py, reading the same per-case workbooks in
data/results_verification/Tables/.

  fig_dc_nameplate      (a) 2050 generation and storage nameplate capacity by
                        technology, summed over all regions, base vs 5 DC cases
                        (stacked bars); (b) per-region % change per technology
  fig_dc_lcoe           (a) aggregate LCOE split into cost components (stacked
                        bars); (b) total LCOE per region + aggregate
  dc_lcoe_comparison.csv  the per-region and total LCOE numbers as a table

The LCOE sheet is a matrix (30 region columns + an "All regions" column, one row
per cost component labelled in column 33); the Nameplate sheet has a
"2050-NAMEPLATE-GW" block of 30 region rows x technology columns.

Usage:
    python -m scripts.plot_dc_economics
    python -m scripts.plot_dc_economics --tables path/to/Tables --outdir figs/
"""
import argparse
import csv
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import numpy as np
import matplotlib.pyplot as plt
import openpyxl

from scripts.plot_style import (
    apply_dc_style as apply_style, DC_STRATEGIES, C_BASELINE,
    add_region_bands, cap_region, minor_ticks, region_sort_key,
    GRID, MUTED, INK, INK_SECONDARY,
)
from scripts.plot_dc_tables import BASE_FILE, CASES, EXCLUDE, _num

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TABLES = REPO_ROOT / "data" / "results_verification" / "Tables"
DEFAULT_OUT = REPO_ROOT / "data" / "results_verification"

apply_style()

# --- Table 3: LCOE cost-component groups (US cents/kWh) -----------------------
# The 15 component rows sum exactly to LCOE-C/KWH-EVERYTHING; here they are
# collapsed into five interpretable buckets for a stacked bar.
LCOE_GROUPS = [
    ("Electricity generation", "#2a78d6",
        ["LCOE-C/KWH-ALL-ELEC-GEN", "LCOE-C/KWH-ADDED-HYDRO-TURBS"]),
    ("Transmission & distribution", "#c98500",
        ["LCOE-C/KWH-SHORT-DIST-TRANS", "LCOE-C/KWH-LONG-DIST-HVDC",
         "LCOE-C/KWH-DISTRIBUTION"]),
    ("Electricity storage", "#1baf7a",
        ["LCOE-C/KWH-LI-BATT-STORAGE", "LCOE-C/KWH-H2-FOR-GRID-ELEC",
         "LCOE-C/KWH-CSP+PHS-STORAGE"]),
    ("Hydrogen (production, compression, storage)", "#4a3aa7",
        ["LCOE-C/KWH-H2-PROD+COMP+STOR"]),
    ("Heat & cold storage", "#e34948",
        ["LCOE-C/KWH-SOLAR+GEO-HEAT", "LCOE-C/KWH-CWSTES+PCMICE-STOR",
         "LCOE-C/KWH-HWSTES-STORAGE", "LCOE-C/KWH-UTES-STORAGE",
         "LCOE-C/KWH-HEAT-PUMPS", "LCOE-C/KWH-HEAT-BRICKS"]),
]
LCOE_TOTAL = "LCOE-C/KWH-EVERYTHING"
ANNUAL_COST = "ANNUAL-ENERGY-COST-$BIL/YR"
LCOE_YLABEL = "2050 levelized cost of all energy (2023 USD cents/kWh)"

# --- Table 2: 2050 nameplate technology groups (GW) --------------------------
NP_TECHS = ["ONWIND", "OFFWIND", "RESPV", "CGOVPV", "UTILPV", "CSP", "GEOELEC",
            "HYDRO", "WAVE", "TIDAL", "SLTHMHEAT", "GEOHEAT", "EGS"]
NP_GROUPS = [
    ("Onshore wind",            "#2a78d6", ["ONWIND"]),
    ("Offshore wind",           "#1baf7a", ["OFFWIND"]),
    ("Utility-scale PV",        "#c98500", ["UTILPV"]),
    ("Rooftop + commercial PV", "#eb6834", ["RESPV", "CGOVPV"]),
    ("Hydropower",              "#4a3aa7", ["HYDRO"]),
    ("Geothermal (incl. EGS)",  "#008300", ["GEOELEC", "EGS"]),
    ("Other (CSP, wave, tidal, solar/geo heat)", "#6e7377",
        ["CSP", "WAVE", "TIDAL", "SLTHMHEAT", "GEOHEAT"]),
]

# Storage nameplate = discharge power (GW) from the ChargeDisch per-region blocks.
# Stacked bottom -> top.  (label, colour, [ChargeDisch keys])
STORAGE_GROUPS = [
    ("Firebrick",     "#b22222", ["HEAT-BRICKS"]),
    ("UTES",          "#8c6d31", ["UTES-HEAT", "UTES-ELEC"]),
    ("HW-STES",       "#e07b39", ["HW-STES"]),
    ("ICE",           "#7fb2d6", ["ICE"]),
    ("CW-STES",       "#1baf7a", ["CW-STES"]),
    ("Grid H2",       "#c51b7d", ["H2-GRID-ELEC"]),
    ("Batteries",     "#333333", ["BATTERIES"]),
    ("CSPS",          "#e9c46a", ["CSP-ELEC", "CSP-PCM"]),
    ("PHS",           "#2a78d6", ["PHS"]),
]
# Per-technology marker + colour (Table 2b small multiples).  Each DC case only
# moves the technologies of its supply strategy, so each panel shows just those.
# Generation nameplate + the two storage capacities that distinguish the cases
# (battery vs hydrogen).  Storage is the discharge-power rating (GW), consistent
# with the generation nameplate axis.
TECH_STYLE = {
    "ONWIND":  ("Onshore wind",     "#2a78d6", "o"),
    "OFFWIND": ("Offshore wind",    "#1baf7a", "s"),
    "UTILPV":  ("Utility-scale PV", "#c98500", "D"),
    "RESPV":   ("Rooftop PV",       "#eb6834", "v"),
    "CGOVPV":  ("Commercial PV",    "#4a3aa7", "^"),
    "EGS":     ("EGS geothermal",   "#008300", "P"),
    "BATT":    ("Battery storage",           "#333333", "X"),
    "H2GRID":  ("Hydrogen storage (grid)",   "#c51b7d", "*"),
}
# legend / stack order (generation first, then storage)
TECH_ORDER = ["ONWIND", "OFFWIND", "UTILPV", "RESPV", "CGOVPV", "EGS", "BATT", "H2GRID"]
NP_VIEW_CAP = 300.0      # per-panel % axis is clipped to a robust range within +/- these
NP_VIEW_FLOOR = -100.0   # bounds; points beyond are drawn as arrows with the value.

# Marker AREA encodes the base (no-DC) capacity of that technology+region (GW),
# log-scaled between these bounds so the ~4-orders-of-magnitude range is legible.
SIZE_MIN, SIZE_MAX = 8.0, 260.0
SIZE_LO_GW, SIZE_HI_GW = 1.0, 15000.0
SIZE_REF_GW = [1, 10, 100, 1000, 10000]   # size-legend reference capacities


def _msize(gw):
    """Base capacity (GW) -> scatter marker area (points^2), log-scaled."""
    v = float(np.clip(gw if gw else SIZE_LO_GW, SIZE_LO_GW, SIZE_HI_GW))
    t = (np.log10(v) - np.log10(SIZE_LO_GW)) / (np.log10(SIZE_HI_GW) - np.log10(SIZE_LO_GW))
    return SIZE_MIN + (SIZE_MAX - SIZE_MIN) * t


def _load(path):
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


# ---- sheet readers -----------------------------------------------------------
def _lcoe(wb):
    """(region_names[30], {LABEL: (per_region{name:val}, all_regions_val)}).

    Keeps the FIRST occurrence of each label (the block carrying per-region cols)."""
    rows = list(wb["LCOE"].iter_rows(values_only=True))
    regions = [str(rows[1][c]).strip() for c in range(1, 31)]
    by = {}
    for r in rows:
        if len(r) > 32 and isinstance(r[32], str) and r[32].strip():
            lab = r[32].strip().upper()
            if lab in by:
                continue
            pr = {regions[c - 1]: _num(r[c]) for c in range(1, 31)}
            by[lab] = (pr, _num(r[31]))
    return regions, by


def _agg_series(wb, exclude_grl_isl):
    """{LABEL: aggregate value} from the LCOE sheet's summary column.

    The default aggregate is the "All regions" column (index 31 / Excel col AF).
    The RBH workbook also carries an "All regions except Greenland & Iceland"
    column (index 34 / Excel col AI); when exclude_grl_isl is set that column is
    read, consistent with excluding GRL/ISL rooftop everywhere else.
    The two columns differ almost entirely in Li-battery storage, because
    GRL/ISL rooftop carries an outsized battery build."""
    rows = list(wb["LCOE"].iter_rows(values_only=True))
    col = 31
    src = "All regions"
    if exclude_grl_isl:
        for r in rows:
            for ci, v in enumerate(r):
                if isinstance(v, str) and "except" in v.lower() and "greenland" in v.lower():
                    col, src = ci, "All regions except Greenland & Iceland"
                    break
            else:
                continue
            break
    out = {}
    for r in rows:
        if len(r) > 32 and isinstance(r[32], str) and r[32].strip():
            lab = r[32].strip().upper()
            v = r[col] if len(r) > col else None
            if lab not in out and isinstance(v, (int, float)):
                out[lab] = float(v)
    return out, src


def _nameplate_2050(wb):
    """{region: {tech: GW}} for the 2050 build-out block (30 region rows)."""
    rows = list(wb["Nameplate"].iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows)
                 if isinstance(r[0], str) and r[0].strip() == "2050-NAMEPLATE-GW")
    out = {}
    for r in rows[start + 1:start + 31]:
        reg = r[0]
        if not isinstance(reg, str) or not reg.strip() or reg.strip().upper().startswith("TOTAL"):
            break
        out[reg.strip()] = {t: (_num(r[ci]) or 0.0) for ci, t in enumerate(NP_TECHS, start=1)}
    return out


def _storage_2050(wb):
    """Ordered list (model region order) of {'BATT': gw, 'H2GRID': gw} from the
    ChargeDisch sheet: battery peak discharge (right block, col 15) and the
    H2-for-grid-electricity discharge (per-region left blocks, H2-GRID-ELEC row).
    Returned as a list so it can be aligned positionally with the Nameplate
    regions (the two sheets spell some regions differently, e.g. Mideast)."""
    rows = list(wb["ChargeDisch"].iter_rows(values_only=True))
    batt = []
    for r in rows[1:]:
        v = r[14] if len(r) > 14 else None
        if isinstance(v, str) and v.strip() and not v.strip().lower().startswith(("all", "total")):
            batt.append(_num(r[15]) or 0.0)
    h2, cur = [], None
    for r in rows:
        c0 = r[0]
        c1 = r[1] if len(r) > 1 else None
        if isinstance(c1, str) and c1.strip() == "PKCHARGE(GW)":     # start of a region block
            cur = {"h": 0.0}
            h2.append(cur)
        if cur is not None and isinstance(c0, str) and c0.strip().upper() == "H2-GRID-ELEC":
            cur["h"] = _num(r[2]) or 0.0                              # col 2 = discharge (GW)
    m = min(len(batt), len(h2))
    return [{"BATT": batt[i], "H2GRID": h2[i]["h"]} for i in range(m)]


def _storage_totals(wb):
    """{ChargeDisch row key: total discharge power (GW) summed over all regions}."""
    rows = list(wb["ChargeDisch"].iter_rows(values_only=True))
    out = {}
    for r in rows:
        c0 = r[0]
        c1 = r[1] if len(r) > 1 else None
        if isinstance(c1, str) and c1.strip() == "PKCHARGE(GW)":     # region-block header
            continue
        if isinstance(c0, str):
            v = _num(r[2])                                            # col 2 = DISCHARGE (GW)
            if v is not None:
                out[c0.strip().upper()] = out.get(c0.strip().upper(), 0.0) + v
    return out


def _stacked_bar(ax, labels, groups, group_vals, totals, ylabel, totfmt, valfmt, width=0.46):
    """Shared stacked-bar renderer (cases on x, one stack segment per group).

    Every segment is labelled with its value: inside the column when the segment
    is tall enough, otherwise to the right of the column with a leader line, and
    thin adjacent segments are nudged apart so their labels don't collide."""
    import matplotlib.patheffects as pe
    n = len(labels)
    xs = np.arange(n)
    ymax = max(totals) * 1.18
    min_gap = 0.043 * ymax
    ax.grid(axis="y", color=GRID, lw=0.7, zorder=0)
    halo = [pe.withStroke(linewidth=1.6, foreground="black")]

    bottoms = np.zeros(n)
    for (glabel, gcolor, _), gv in zip(groups, group_vals):
        gv = np.asarray(gv, dtype=float)
        ax.bar(xs, gv, width=width, bottom=bottoms, color=gcolor, label=glabel,
               edgecolor="white", linewidth=0.6, zorder=3)
        bottoms += gv

    for ci in range(n):                       # value labels, one column at a time
        outside, b = [], 0.0
        for gi, (_gl, gcolor, _) in enumerate(groups):
            v = float(group_vals[gi][ci])
            if v <= 0:
                continue
            yc = b + v / 2.0
            if v / ymax >= 0.05:              # fits inside the segment
                ax.text(ci, yc, valfmt.format(v), ha="center", va="center", fontsize=8,
                        color="white", zorder=6, path_effects=halo)
            else:                             # too thin: collect for outside placement
                outside.append([yc, valfmt.format(v), gcolor])
            b += v
        outside.sort(key=lambda e: e[0])      # push labels apart, bottom-up
        last = -1e9
        for yc, text, gcolor in outside:
            y = max(yc, last + min_gap)
            last = y
            xseg, xlab = ci + width / 2, ci + width / 2 + 0.07
            ax.plot([xseg, xlab], [yc, y], color=gcolor, lw=0.6, zorder=5)
            ax.text(xlab + 0.02, y, text, ha="left", va="center", fontsize=7.5,
                    color=gcolor, zorder=6)
    for x, t in zip(xs, totals):
        ax.text(x, t, totfmt.format(t), ha="center", va="bottom", fontsize=10.5,
                fontweight="bold", color=INK)
    ax.set_ylim(0, ymax)
    ax.set_ylabel(ylabel)
    ax.set_xticks(xs)
    ax.set_xticklabels(labels, fontsize=10)
    ax.set_xlim(-0.6, n - 0.4)
    ax.tick_params(axis="x", length=0)
    minor_ticks(ax)
    ax.legend(loc="upper left", bbox_to_anchor=(1.01, 1.0), fontsize=9.5,
              frameon=False, borderaxespad=0.0)


# ---- Table 3 : LCOE breakdown (panel a of the combined LCOE figure) ----------
def _draw_lcoe_breakdown(ax, tables):
    labels, totals, srcs = [], [], []
    group_vals = [[] for _ in LCOE_GROUPS]
    for fn, label, _c in CASES:
        wb = _load(tables / fn)
        agg, src = _agg_series(wb, fn in EXCLUDE)
        wb.close()
        labels.append(label)
        srcs.append(src)
        for gi, (_gl, _gc, comps) in enumerate(LCOE_GROUPS):
            group_vals[gi].append(sum(agg.get(c.upper(), 0.0) for c in comps))
        totals.append(agg.get(LCOE_TOTAL, 0.0))

    _stacked_bar(ax, labels, LCOE_GROUPS, group_vals, totals, LCOE_YLABEL, "{:.2f}", "{:.2f}")
    print("    total LCOE (cents/kWh): "
          + "; ".join(f"{l} {t:.2f}" for l, t in zip(labels, totals)))
    for l, s in zip(labels, srcs):
        if "except" in s:
            print(f"    {l} aggregate uses '{s}'")


# ---- Table 2a : global nameplate, generation + storage columns (panel a) -----
def _draw_nameplate_global(ax, legend_ax, tables):
    """Two stacked columns per case: generation (left) and storage (right),
    both in TW.  Storage discharge power comes from the ChargeDisch sheet."""
    import matplotlib.patheffects as pe
    labels = [l for _, l, _ in CASES]
    gen_vals = [[] for _ in NP_GROUPS]
    sto_vals = [[] for _ in STORAGE_GROUPS]
    gen_tot, sto_tot = [], []
    for fn, _label, _c in CASES:
        wb = _load(tables / fn)
        data = _nameplate_2050(wb)
        stot = _storage_totals(wb)
        wb.close()
        excl = EXCLUDE.get(fn, set())
        techtot = {t: 0.0 for t in NP_TECHS}
        for reg, tv in data.items():
            if reg in excl:
                continue
            for t in NP_TECHS:
                techtot[t] += tv[t]
        for gi, (_gl, _gc, techs) in enumerate(NP_GROUPS):
            gen_vals[gi].append(sum(techtot[t] for t in techs) / 1000.0)
        gen_tot.append(sum(techtot.values()) / 1000.0)
        for gi, (_gl, _gc, keys) in enumerate(STORAGE_GROUPS):
            sto_vals[gi].append(sum(stot.get(k, 0.0) for k in keys) / 1000.0)
        sto_tot.append(sum(sto_vals[gi][-1] for gi in range(len(STORAGE_GROUPS))))

    n = len(labels)
    xs = np.arange(n)
    dx, w = 0.23, 0.38
    ymax = max(max(gen_tot), max(sto_tot)) * 1.16
    min_gap = 0.045 * ymax
    halo = [pe.withStroke(linewidth=1.6, foreground="black")]
    ax.grid(axis="y", color=GRID, lw=0.7, zorder=0)

    def stack(xc, groups, vals, tot, side):
        """Draw one stacked column type; label big segments inside and smaller
        ones (>=0.3 TW) outside with a leader line on the given side (generation
        labels to the left, storage to the right) -- de-overlapped vertically."""
        bottoms = np.zeros(n)
        outside = [[] for _ in range(n)]
        for (glabel, gcolor, _), gv in zip(groups, vals):
            gv = np.asarray(gv, float)
            ax.bar(xc, gv, width=w, bottom=bottoms, color=gcolor, label=glabel,
                   edgecolor="white", linewidth=0.5, zorder=3)
            for ci, (x, v, b) in enumerate(zip(xc, gv, bottoms)):
                if v <= 0:
                    continue
                yc = b + v / 2
                if v / ymax >= 0.028:
                    ax.text(x, yc, f"{v:.1f}", ha="center", va="center", fontsize=7,
                            color="white", zorder=6, path_effects=halo)
                elif v >= 0.3:
                    outside[ci].append([yc, f"{v:.1f}", gcolor])
            bottoms += gv
        for ci, items in enumerate(outside):
            items.sort(key=lambda e: e[0])
            last = -1e9
            for yc, text, gcolor in items:
                y = max(yc, last + min_gap)
                last = y
                x = xc[ci]
                if side == "left":
                    xseg, xlab = x - w / 2, x - w / 2 - 0.05
                    ax.plot([xseg, xlab], [yc, y], color=gcolor, lw=0.6, zorder=5)
                    ax.text(xlab - 0.02, y, text, ha="right", va="center", fontsize=7,
                            color=gcolor, zorder=6)
                else:
                    xseg, xlab = x + w / 2, x + w / 2 + 0.05
                    ax.plot([xseg, xlab], [yc, y], color=gcolor, lw=0.6, zorder=5)
                    ax.text(xlab + 0.02, y, text, ha="left", va="center", fontsize=7,
                            color=gcolor, zorder=6)
        for x, t in zip(xc, tot):
            ax.text(x, t, f"{t:.1f}", ha="center", va="bottom", fontsize=8.5,
                    fontweight="bold", color=INK)

    stack(xs - dx, NP_GROUPS, gen_vals, gen_tot, "left")
    stack(xs + dx, STORAGE_GROUPS, sto_vals, sto_tot, "right")

    # two-level x labels: Gen/Stor under each bar, case name under the pair
    tr = ax.get_xaxis_transform()
    for x in xs:
        ax.text(x - dx, -0.02, "Gen", transform=tr, ha="center", va="top", fontsize=7.5,
                color=INK_SECONDARY)
        ax.text(x + dx, -0.02, "Stor", transform=tr, ha="center", va="top", fontsize=7.5,
                color=INK_SECONDARY)
    ax.set_ylim(0, ymax)
    ax.set_ylabel("Nameplate capacity needed in 2050 (TW)")
    ax.set_xticks(xs)
    ax.set_xticklabels(labels, fontsize=10)
    ax.tick_params(axis="x", length=0, pad=16)
    ax.set_xlim(-0.85, n - 0.15)
    minor_ticks(ax)

    # legends (generation + storage) in the side panel, both flush-left
    from matplotlib.patches import Patch
    legend_ax.axis("off")
    gh = [Patch(fc=c, label=l) for l, c, _ in NP_GROUPS]
    sh = [Patch(fc=c, label=l) for l, c, _ in STORAGE_GROUPS]
    leg1 = legend_ax.legend(handles=gh, loc="upper left", bbox_to_anchor=(0.0, 1.0),
                            fontsize=9, frameon=False, title="Generation",
                            title_fontsize=10.5, alignment="left", borderaxespad=0.0)
    legend_ax.add_artist(leg1)
    legend_ax.legend(handles=sh, loc="lower left", bbox_to_anchor=(0.0, 0.0),
                     fontsize=9, frameon=False, title="Storage (discharge power)",
                     title_fontsize=10.5, alignment="left", borderaxespad=0.0)
    print("    2050 nameplate gen/storage (TW): "
          + "; ".join(f"{l} {g:.1f}/{s:.1f}" for l, g, s in zip(labels, gen_tot, sto_tot)))


# ---- Table 2b : per-region nameplate % change, BY TECHNOLOGY ------------------
def _robust_view(vals):
    """Per-panel y-limits clipped to a robust range within the hard bounds."""
    fin = vals[np.isfinite(vals)]
    if not fin.size:
        return -1.0, 10.0
    lo = max(NP_VIEW_FLOOR, min(0.0, float(np.percentile(fin, 3))))
    hi = min(NP_VIEW_CAP, float(np.percentile(fin, 97)))
    if hi <= lo:
        hi = lo + 10.0
    pad = 0.12 * (hi - lo)
    return lo - 0.35 * pad, hi + pad


def _combined_2050(wb):
    """{region: {generation techs ... , BATT, H2GRID}} in GW, merging the
    Nameplate generation block with the ChargeDisch storage discharge power
    (aligned positionally, since the two sheets spell some regions differently)."""
    gen = _nameplate_2050(wb)
    stor = _storage_2050(wb)
    for i, reg in enumerate(gen):
        if i < len(stor):
            gen[reg].update(stor[i])
        else:
            gen[reg].update({"BATT": 0.0, "H2GRID": 0.0})
    return gen


def _draw_nameplate_regional(axes, tables):
    """Draw the per-region % nameplate change small multiples into `axes`
    (>= len(CASES) axes; the last is used for the shared legend)."""
    wb = _load(tables / BASE_FILE)
    base = _combined_2050(wb)
    wb.close()
    regions = list(base.keys())
    order = sorted(regions, key=region_sort_key)   # alphabetical, shared across all figures
    n = len(order)
    xs = np.arange(n)

    # per case: which technologies (generation + storage) actually change, and
    # their per-region % change.  |global delta|>1 GW selects exactly the
    # strategy's technologies (incl. battery for WSB, hydrogen for WSH, both for
    # WSBH/RBH).
    panels, skipped = [], []
    for fn, label, _c in CASES[1:]:
        wb = _load(tables / fn)
        cur = _combined_2050(wb)
        wb.close()
        excl = EXCLUDE.get(fn, set())
        changed = [t for t in TECH_ORDER
                   if abs(sum(cur[r].get(t, 0.0) for r in regions)
                          - sum(base[r].get(t, 0.0) for r in regions)) > 1.0]
        series = {}
        for t in changed:
            vals = np.array([
                ((cur[r][t] - base[r][t]) / base[r][t] * 100.0)
                if (base[r].get(t, 0.0) > 1e-6 and r not in excl) else np.nan
                for r in order])
            series[t] = vals
            n_zero = sum(1 for r in order if base[r].get(t, 0.0) <= 1e-6 and r not in excl)
            if n_zero:
                skipped.append((label.replace("\n", " "), TECH_STYLE[t][0], n_zero))
        panels.append((label.replace("\n", " "), series))

    offscale, shown = [], set()
    for pi, (label, series) in enumerate(panels):
        ax = axes[pi]
        add_region_bands(ax, n)
        ax.axhline(0.0, color=MUTED, lw=0.8, zorder=1)
        allv = (np.concatenate([v[np.isfinite(v)] for v in series.values()])
                if series else np.array([0.0]))
        ylo, yhi = _robust_view(allv)
        off = []   # (case, tech, region, value) beyond the axis -> arrow only
        for t, vals in series.items():
            shown.add(t)
            lbl, color, marker = TECH_STYLE[t]
            sizes = np.array([_msize(base[r].get(t, 0.0)) for r in order])   # marker area = base GW
            onscale = np.where((vals >= ylo) & (vals <= yhi), vals, np.nan)
            ax.scatter(xs, onscale, s=sizes, color=color, marker=marker, label=lbl,
                       edgecolors="white", linewidths=0.5, zorder=4)
            # off-scale points get a bare arrow (values printed to stdout so the
            # crowded high tails of the rooftop case stay legible).
            for xi, v, reg in zip(xs, vals, order):
                if np.isfinite(v) and v > yhi:
                    ax.text(xi, yhi, "$\\uparrow$", ha="center", va="top",
                            fontsize=8, color=color, zorder=5)
                    off.append((label, lbl, reg, v))
                elif np.isfinite(v) and v < ylo:
                    ax.text(xi, ylo, "$\\downarrow$", ha="center", va="bottom",
                            fontsize=8, color=color, zorder=5)
                    off.append((label, lbl, reg, v))
        offscale.extend(off)
        ax.set_ylim(ylo, yhi)
        ax.set_xlim(-0.7, n - 0.3)
        ax.text(0.0, 1.015, label, transform=ax.transAxes, fontweight="bold",
                fontsize=12, va="bottom")
        ax.set_xticks(range(n))
        ax.set_xticklabels([cap_region(r) for r in order], rotation=90, fontsize=6.4)
        ax.tick_params(axis="x", length=0)
        minor_ticks(ax)
        if pi % 3 == 0:
            ax.set_ylabel("Change in each technology's nameplate\n"
                          "capacity vs base (no DC) (%)")

    # shared legends in the empty (bottom-right) panel: technology (colour+marker)
    # and base-capacity (marker size)
    from matplotlib.lines import Line2D
    tech_handles = [Line2D([0], [0], marker=TECH_STYLE[t][2], linestyle="none",
                           markerfacecolor=TECH_STYLE[t][1], markeredgecolor="white",
                           markersize=13, label=TECH_STYLE[t][0])
                    for t in TECH_ORDER if t in shown]
    size_handles = [Line2D([0], [0], marker="o", linestyle="none", markerfacecolor="#8a8f96",
                           markeredgecolor="white", markersize=np.sqrt(_msize(v)),
                           label=f"{v:,} GW")
                    for v in SIZE_REF_GW]
    for j in range(len(panels), len(axes)):
        axes[j].axis("off")
    lax = axes[len(axes) - 1]
    leg_t = lax.legend(handles=tech_handles, loc="center left", bbox_to_anchor=(0.0, 0.5),
                       fontsize=12.5, frameon=False, title="Technology", title_fontsize=14,
                       labelspacing=0.85, handletextpad=0.6)
    lax.add_artist(leg_t)
    lax.legend(handles=size_handles, loc="center right", bbox_to_anchor=(1.0, 0.5),
               fontsize=12, frameon=False, title="Base (no-DC) capacity",
               title_fontsize=13, labelspacing=1.3, handletextpad=0.8, borderpad=1.0)

    for label, series in panels:
        techs = ", ".join(TECH_STYLE[t][0] for t in series)
        print(f"    {label}: technologies shown = {techs}")
    if offscale:
        note = "; ".join(f"{c} {t} {reg} {v:+,.0f}%" for c, t, reg, v in offscale)
        print(f"    off-scale (arrow only): {note}")
    if skipped:
        note = "; ".join(f"{c} {t} ({z} zero-base region{'s' if z > 1 else ''})"
                         for c, t, z in skipped)
        print(f"    zero-base regions omitted (no % defined): {note}")


# ---- Table 1 : LCOE per region + aggregate -----------------------------------
def _gather_lcoe(tables):
    """{case_label: (per_region_lcoe{}, all_regions_lcoe)}, region_names, base cost."""
    case_vals = {}
    region_names, base_cost = None, {}
    for fn, label, _c in CASES:
        wb = _load(tables / fn)
        regs, by = _lcoe(wb)
        agg, _src = _agg_series(wb, fn in EXCLUDE)   # excl. GRL/ISL aggregate for RBH
        wb.close()
        region_names = regs
        pr, _allv = by.get(LCOE_TOTAL, ({}, None))
        excl = EXCLUDE.get(fn, ())
        pr = {k: (np.nan if k in excl else v) for k, v in pr.items()}
        case_vals[label] = (pr, agg.get(LCOE_TOTAL))
        if fn == BASE_FILE:
            base_cost = by.get(ANNUAL_COST, ({}, None))[0]
    return case_vals, region_names, base_cost


def write_lcoe_table(case_vals, region_names, outdir):
    labels = [l for _, l, _ in CASES]
    with open(outdir / "dc_lcoe_comparison.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Region"] + [l.replace("\n", " ") + " LCOE (US cents/kWh)" for l in labels])
        for reg in region_names:
            row = [reg]
            for l in labels:
                v = case_vals[l][0].get(reg)
                row.append("" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{v:.3f}")
            w.writerow(row)
        agg = ["All regions"]
        for l in labels:
            v = case_vals[l][1]
            agg.append("" if v is None else f"{v:.3f}")
        w.writerow(agg)
    print("  wrote dc_lcoe_comparison.csv")


def _draw_lcoe_by_region(ax, case_vals, region_names):
    order = sorted(region_names, key=region_sort_key)   # alphabetical, shared across all figures
    n = len(order)
    xs = np.arange(n)
    gap = 1.6
    xagg = n - 1 + gap + 1     # aggregate column sits to the right of a separator

    add_region_bands(ax, n)
    ax.axvline((n - 1 + xagg) / 2.0, color=MUTED, lw=0.9, ls=(0, (2, 2)), zorder=1)

    def draw(xpos, getval):
        base_pr = case_vals[CASES[0][1]][0]
        ax.scatter(xpos, getval(base_pr), s=70, facecolors="none",
                   edgecolors=C_BASELINE, linewidths=1.5, marker="o",
                   label="Base (no DC)", zorder=3)
        for (_fn, label, _c), (_key, dclab, color, marker) in zip(CASES[1:], DC_STRATEGIES):
            ax.scatter(xpos, getval(case_vals[label][0]), s=40, color=color, marker=marker,
                       edgecolors="white", linewidths=0.6, label=dclab, zorder=4)

    draw(xs, lambda pr: [pr.get(r) if not (isinstance(pr.get(r), float) and np.isnan(pr.get(r)))
                         else np.nan for r in order])
    # aggregate ("All regions") column: getval ignores per-region dict, uses the [1] slot
    base_all = case_vals[CASES[0][1]][1]
    ax.scatter([xagg], [base_all], s=70, facecolors="none", edgecolors=C_BASELINE,
               linewidths=1.5, marker="o", zorder=3)
    for (_fn, label, _c), (_key, _dclab, color, marker) in zip(CASES[1:], DC_STRATEGIES):
        ax.scatter([xagg], [case_vals[label][1]], s=40, color=color, marker=marker,
                   edgecolors="white", linewidths=0.6, zorder=4)

    ax.set_ylabel(LCOE_YLABEL)
    ax.legend(loc="upper left", ncols=6, handletextpad=0.2, columnspacing=1.0,
              borderaxespad=0.3, bbox_to_anchor=(0.0, 1.10))
    xticks = list(range(n)) + [xagg]
    xlabels = [cap_region(r) for r in order] + ["All regions"]
    ax.set_xticks(xticks)
    ax.set_xticklabels(xlabels, rotation=45, ha="right", rotation_mode="anchor")
    ax.set_xlim(-0.7, xagg + 0.7)
    ax.tick_params(axis="x", length=0)
    minor_ticks(ax)


def _panel_label(ax, text, y=1.045):
    ax.text(0.0, y, text, transform=ax.transAxes, fontweight="bold",
            fontsize=16, va="bottom", ha="left")


# ---- combined figures (a = global, b = regional) -----------------------------
def fig_dc_nameplate(tables, outdir):
    """Combined nameplate figure: (a) global generation + storage columns,
    (b) per-region % change small multiples."""
    fig = plt.figure(figsize=(20.0, 16.0))
    gs = fig.add_gridspec(3, 3, height_ratios=[1.15, 1.0, 1.0],
                          hspace=0.42, wspace=0.22)
    axa = fig.add_subplot(gs[0, :2])
    lax = fig.add_subplot(gs[0, 2])
    _draw_nameplate_global(axa, lax, tables)
    _panel_label(axa, "a")
    reg_axes = [fig.add_subplot(gs[1 + r, c]) for r in range(2) for c in range(3)]
    _draw_nameplate_regional(reg_axes, tables)
    _panel_label(reg_axes[0], "b", y=1.14)   # lift clear of the "DC EGS" panel title
    for ext in ("pdf", "png"):
        fig.savefig(outdir / f"fig_dc_nameplate.{ext}")
    plt.close(fig)
    print("  wrote fig_dc_nameplate.pdf/.png")


def fig_dc_lcoe(tables, outdir):
    """Combined LCOE figure: (a) global cost-component breakdown,
    (b) total LCOE per region + aggregate."""
    case_vals, region_names, _bc = _gather_lcoe(tables)
    write_lcoe_table(case_vals, region_names, outdir)
    fig = plt.figure(figsize=(16.0, 12.5))
    gs = fig.add_gridspec(2, 1, height_ratios=[1.0, 1.1], hspace=0.30)
    axa = fig.add_subplot(gs[0])
    _draw_lcoe_breakdown(axa, tables)
    _panel_label(axa, "a")
    axb = fig.add_subplot(gs[1])
    _draw_lcoe_by_region(axb, case_vals, region_names)
    _panel_label(axb, "b")
    for ext in ("pdf", "png"):
        fig.savefig(outdir / f"fig_dc_lcoe.{ext}")
    plt.close(fig)
    print("  wrote fig_dc_lcoe.pdf/.png")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tables", type=Path, default=DEFAULT_TABLES)
    ap.add_argument("--outdir", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args(argv)
    args.outdir.mkdir(parents=True, exist_ok=True)

    fig_dc_nameplate(args.tables, args.outdir)   # Table 2 (a + b)
    fig_dc_lcoe(args.tables, args.outdir)         # Tables 1 & 3 (a + b) + CSV


if __name__ == "__main__":
    main()
