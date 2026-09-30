"""Structural-difference figures: what the optimisers actually changed.

Companion to plot_comparison.py.  Where those figures show the OUTCOME (cost,
land, solve time), these show the STRUCTURE behind it - which design levers moved
and how the system was rebuilt.  Both optimised cases are compared against the
trial-and-error baseline:

    GA (from trial-and-error)   - blue
    GA (from scratch, extended) - orange

Data comes from the two already-exported artefacts:

  data/results_verification/comparison_summary.csv   (region order = system cost)
  data/results_verification/results_export.xlsx      (per region x case detail)

Figures:

  fig_factor_changes            per optimised factor, the distribution across
                                regions of the ratio optimised / trial-and-error
                                (log axis; range bar per case: full range, IQR,
                                median) plus switched-on/off counts.
  fig_cost_change_by_category   per region, the change in cost (2023 c/kWh) vs trial-
                                and-error decomposed into five additive cost groups
                                (diverging stacked bar: increases up, savings down)
                                with the net change marked; (a) GA-from-trial-and-
                                error, (b) GA-from-scratch, each on its own y-scale.
  fig_generation_mix_change     wind and solar share of TOTAL generation per region
                                (two panels), dumbbell of trial-and-error vs both
                                optimised cases; the connector is the change.

A region is absent where the baseline run is infeasible, so no change can be
defined (e.g. a region whose trial-and-error re-run did not converge).  Same
formatting as the other figures (landscape, regions on x, larger fonts,
colour-blind-safe palette, no on-figure titles; caption numbers to stdout).

Usage:
    python -m scripts.plot_structure
    python -m scripts.plot_structure --case-a "GA (trial-error)" --case-b "GA (scratch2)"
"""
import argparse
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.ticker import NullLocator
import openpyxl
import pandas as pd

from scripts.plot_style import (
    apply_style, region_order, add_region_bands, style_region_axis, minor_ticks,
    region_label,
    GRID, MUTED, INK, INK_SECONDARY, C_BASELINE, C_GA, C_SCRATCH,
    LABEL_BASELINE, LABEL_GA, LABEL_SCRATCH, TAE_IT, plain_label,
)
from scripts.plot_comparison import _dumbbell, SZ_BASE, SZ_GA, SZ_SCRATCH

REPO_ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = REPO_ROOT / "data" / "results_verification"
DEFAULT_CSV = RESULTS_DIR / "comparison_summary.csv"
DEFAULT_XLSX = RESULTS_DIR / "results_export.xlsx"

BASELINE_CASE = "Baseline"
DEFAULT_CASE_A = "GA (trial-error)"
DEFAULT_CASE_B = "GA (scratch2)"

G_FACTORS = "Optimised factors"
G_COST = "Cost by category (c/kWh, MN)"
G_GEN = "Generation (TWh/yr)"
GEN_EXCLUDE = {"Total supply", "Total gen (sum)"}
# Generation sources shown in fig_generation_mix_change: the two the GA moves
# (via the wind / PV capacity factors).  Solar heat is also GA-adjusted (FACSHT)
# but is ~1% of generation and barely moves, so it is reported to stdout only;
# hydro, geothermal, wave, tidal are not design variables.
GEN_SHOWN = ["Wind", "Solar"]

# Optimised factors that scale GENERATION capacity (wind / PV / solar-thermal /
# CSP turbine); everything else is storage duration/power + dispatch.  Used to
# group the rows of fig_factor_changes (generation on top).
GEN_FACTORS = {"FACONWIN", "FACOFFWIN", "FACUTILPV", "FACRESPV", "FACCOMPV",
               "FACSHT", "CSPTURBFAC"}
# Factors that are held fixed (not GA design variables) and so are not shown as a
# structural "change": CPERFORM (heat-pump COP) is a physical constant.
FACTOR_EXCLUDE = {"CPERFORM"}

# Human-readable tick labels (the raw keys are model-internal variable names).
FACTOR_LABELS = {
    "FACONWIN": "Onshore wind", "FACOFFWIN": "Offshore wind",
    "FACUTILPV": "Utility-scale PV", "FACRESPV": "Residential PV",
    "FACCOMPV": "Commercial PV", "FACSHT": "Solar thermal heat",
    "CSPTURBFAC": "CSP turbine",
    "BATDISCH": "Battery discharge rate", "STORHBAT": "Battery duration",
    "STORHCOLD": "Cold-storage duration", "STORHHWAT": "Hot-water STES duration",
    "STORHPHS": "Pumped-hydro duration", "STORUGDYS": "Seasonal UTES duration",
    "DAYH2STOR": "H$_2$ storage duration",
    "FCDISCH": "Fuel-cell discharge rate", "FCCHARG": "Electrolyser charge rate",
}

# Additive cost groups for fig_cost_change_by_category, stacked bottom -> top.
# The 15 Fortran c/kWh categories are disjoint and sum to the total;
# ALL-NONH2-STORAGE is the Fortran's own sum of the non-H2 storage rows, so it is
# left out (including it would double-count storage).  Colours: data-viz
# categorical slots 3-7 (blue/orange stay reserved for the two optimised
# cases); validated all-pairs on white, worst pair aqua<->magenta CVD dE 6.1
# (floor band), carried by the 2 px white gaps between segments and by never
# placing the two next to each other in the stack order.
COST_GROUPS = [
    ("Electricity generation", ["ELECTRICITY GEN ONLY"], "#eda100"),
    ("Transmission & distribution",
     ["SHORT-DIST TRANSMISS", "LONG-DIST-TRANS", "DISTRIBUTION"], "#4a3aa7"),
    ("Li-ion battery storage", ["LI-BATTERY STORAGE"], "#1baf7a"),
    ("Hydrogen (grid + non-grid)",
     ["H2-ELEC PROD/STOR/FC", "H2-PROD/COMP/STOR"], "#008300"),
    ("Thermal & other storage",
     ["UTES STORAGE", "HWSTES STORAGE", "CWSTES+PCMICE STOR",
      "HPUMPS FOR HWST+UTES", "INDUS FIREBRICK STOR", "CSPPCM + PHS STORAGE",
      "ADDED-HYDRO-TURBS", "SOLAR+GEOTHERM HEAT"], "#e87ba4"),
]
COST_AGGREGATES = {"ALL-NONH2-STORAGE"}

apply_style()


def load_structure(xlsx_path: Path) -> pd.DataFrame:
    """Flatten the two-row-header 'Results' sheet into a tidy frame whose columns
    are 'Group :: Name' and rows are one region x case each."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    groups, names = list(rows[0]), list(rows[1])
    cols, g = [], None
    for grp, nm in zip(groups, names):
        if grp not in (None, ""):
            g = grp
        cols.append(f"{g} :: {nm}")
    df = pd.DataFrame(rows[2:], columns=cols)
    return df[df["Identification :: Region"].notna()].copy()


def _cols_in_group(df, group, exclude=()):
    return [c for c in df.columns
            if c.startswith(group + " :: ") and c.split(" :: ", 1)[1] not in exclude]


def _case_matrix(df, cols, case, regions):
    sub = (df[df["Identification :: Case"] == case]
           .set_index("Identification :: Region").reindex(regions))
    return sub[cols].apply(pd.to_numeric, errors="coerce")


def _short(col):
    return col.split(" :: ", 1)[1]


def _panel_tag(ax, letter, label):
    """Bold panel letter plus the case it shows, above the top-left corner."""
    ax.text(0.0, 1.02, letter, transform=ax.transAxes, fontweight="bold",
            fontsize=14, va="bottom")
    ax.text(0.022, 1.02, label, transform=ax.transAxes, fontsize=12, va="bottom")


def _stacked_change(ax, xs, parts, width=0.62):
    """Diverging stacked bars: per x, positive parts stack up from 0 and negative
    parts down from 0, in the given order.  `parts` = [(values, colour), ...].
    Returns the (top, bottom) extent of each column."""
    xs = np.asarray(xs)
    top = np.zeros(len(xs))
    bot = np.zeros(len(xs))
    for vals, color in parts:
        v = np.nan_to_num(np.asarray(vals, float))
        for sign, base in ((1, top), (-1, bot)):
            m = v * sign > 0
            if m.any():
                ax.bar(xs[m], v[m], width, bottom=base[m], color=color,
                       edgecolor="white", linewidth=0.8, zorder=3)
                base[m] += v[m]
    return top, bot


def _outlier_cap(extent, ratio=1.5, pad=1.12):
    """Axis limit that keeps the bulk legible when one or two columns dwarf the
    rest: if the largest extent is > `ratio` x the third largest, cap at `pad` x
    the third largest (the clipped columns are annotated); else None."""
    e = np.sort(np.abs(extent[np.isfinite(extent)]))[::-1]
    if len(e) >= 3 and e[0] > ratio * e[2]:
        return pad * e[2]
    return None


def fig_factor_changes(df, outdir, regions, case_a, case_b):
    """Per optimised factor, the distribution across regions of the ratio
    optimised / trial-and-error on a log axis (so halving and doubling sit
    symmetrically about x1).  Each case is a range bar: thin line = full range,
    thick bar = interquartile range, dot = median.  Regions whose trial-and-
    error value is 0 (ratio undefined) or whose optimum is 0 are counted in a
    right-hand column instead ("on" = switched on from 0, "off" = set to 0)."""
    cols = [c for c in _cols_in_group(df, G_FACTORS) if _short(c) not in FACTOR_EXCLUDE]
    base = _case_matrix(df, cols, BASELINE_CASE, regions)
    opt = {case: _case_matrix(df, cols, case, regions) for case in (case_a, case_b)}
    eps = 1e-9

    def _ratios(case, c):
        b, o = base[c], opt[case][c]
        m = (b > eps) & (o > eps)
        return ((o[m] / b[m]).to_numpy(),
                int(((b.abs() <= eps) & (o > eps)).sum()),
                int(((b > eps) & (o.abs() <= eps)).sum()))

    stats = {(case, c): _ratios(case, c) for case in (case_a, case_b) for c in cols}

    def _spread(c):   # larger |log median| of the two cases, for row order
        return max(abs(np.log(np.median(stats[(k, c)][0]))) if len(stats[(k, c)][0]) else 0
                   for k in (case_a, case_b))

    groups = [("Generation capacity",
               sorted([c for c in cols if _short(c) in GEN_FACTORS], key=lambda c: -_spread(c))),
              ("Storage & dispatch",
               sorted([c for c in cols if _short(c) not in GEN_FACTORS], key=lambda c: -_spread(c)))]
    # y layout top -> bottom: a header line per group, then one row per factor
    ypos, headers, y = {}, [], 0.0
    for title, members in groups:
        if not members:
            continue
        headers.append((title, y))
        y -= 0.9
        for c in members:
            ypos[c] = y
            y -= 1.0
        y -= 0.3
    ymin = y + 0.3

    allr = np.concatenate([v[0] for v in stats.values() if len(v[0])])
    lo = 2.0 ** np.floor(np.log2(allr.min()) - 0.5)
    hi = 2.0 ** np.ceil(np.log2(allr.max()) + 0.5)

    fig, ax = plt.subplots(figsize=(12.0, 0.46 * len(ypos) + 2.4))
    ax.set_xscale("log")
    ax.set_xlim(lo, hi)
    ticks = [t for t in (0.01, 0.1, 0.25, 0.5, 1, 2, 4, 10, 100, 1000) if lo <= t <= hi]
    ax.set_xticks(ticks)
    ax.set_xticklabels([f"\u00d7{t:g}" for t in ticks])
    ax.xaxis.set_minor_locator(NullLocator())
    ax.grid(axis="x", color=GRID, lw=0.7, zorder=0)
    ax.axvline(1.0, color=MUTED, lw=1.0, zorder=1)
    ytr = ax.get_yaxis_transform()
    for c, yc in ypos.items():
        for case, color, dy in ((case_a, C_GA, 0.18), (case_b, C_SCRATCH, -0.18)):
            r, n_on, n_off = stats[(case, c)]
            yy = yc + dy
            if len(r):
                q1, med, q3 = np.percentile(r, [25, 50, 75])
                ax.plot([r.min(), r.max()], [yy, yy], color=color, lw=1.3,
                        solid_capstyle="round", zorder=3)
                ax.plot([q1, q3], [yy, yy], color=color, lw=7.5,
                        solid_capstyle="butt", zorder=4)
                ax.scatter(med, yy, s=34, color="white", edgecolors=color,
                           linewidths=1.6, zorder=5)
            cnt = "  ".join(t for t in (f"+{n_on} on" if n_on else "",
                                        f"{n_off} off" if n_off else "") if t)
            if cnt:
                ax.text(1.01, yy, cnt, transform=ytr, ha="left", va="center",
                        fontsize=9, color=INK_SECONDARY)
    for title, yh in headers:
        ax.text(-0.005, yh, title, transform=ytr, ha="right", va="center",
                fontsize=12, fontweight="bold", color=INK)
    ax.text(1.01, headers[0][1], "Switched on / off\n(regions)", transform=ytr,
            ha="left", va="center", fontsize=9, color=INK_SECONDARY)
    ax.set_yticks(list(ypos.values()))
    ax.set_yticklabels([FACTOR_LABELS.get(_short(c), _short(c)) for c in ypos])
    ax.tick_params(axis="y", length=0)
    ax.set_ylim(ymin - 0.2, headers[0][1] + 0.5)
    ax.spines["left"].set_visible(False)
    ax.set_xlabel(f"Optimised factor relative to LOADMATCH ({TAE_IT}) "
                  "(ratio, log scale)")
    handles = [plt.Line2D([], [], color=C_GA, lw=7.5, label=LABEL_GA),
               plt.Line2D([], [], color=C_SCRATCH, lw=7.5, label=LABEL_SCRATCH),
               plt.Line2D([], [], marker="o", ls="", mfc="white", mec=MUTED, mew=1.6,
                          ms=6.5, label="Median"),
               plt.Line2D([], [], color=MUTED, lw=7.5, label="Middle 50% of regions"),
               plt.Line2D([], [], color=MUTED, lw=1.3, label="Full range")]
    ax.legend(handles=handles, loc="lower center", bbox_to_anchor=(0.5, 1.0),
              ncols=5, handlelength=1.6, columnspacing=1.3, borderaxespad=0.3,
              fontsize=10)
    for ext in ("pdf", "png"):
        fig.savefig(outdir / f"fig_factor_changes.{ext}")
    plt.close(fig)

    print("  wrote fig_factor_changes.pdf/.png")
    for case, label in ((case_a, LABEL_GA), (case_b, LABEL_SCRATCH)):
        print(f"    caption stats, {plain_label(label)} (ratio median [IQR] range; n, on, off):")
        for c in ypos:
            r, n_on, n_off = stats[(case, c)]
            if len(r):
                q1, med, q3 = np.percentile(r, [25, 50, 75])
                print(f"      {FACTOR_LABELS.get(_short(c), _short(c)):28s} x{med:.2f} "
                      f"[{q1:.2f}-{q3:.2f}] {r.min():.3g}-{r.max():.3g}; n={len(r)}, "
                      f"on={n_on}, off={n_off}, unchanged(+-0.5%)="
                      f"{int((np.abs(r - 1) < 0.005).sum())}")


def fig_cost_change_by_category(df, outdir, regions, case_a, case_b):
    cols = _cols_in_group(df, G_COST, exclude=COST_AGGREGATES)
    have = {_short(c) for c in cols}
    base = _case_matrix(df, cols, BASELINE_CASE, regions)
    base.columns = [_short(c) for c in cols]
    regs = [r for r in regions if base.loc[r].notna().any()]   # feasible baseline

    def _grouped(case):
        opt = _case_matrix(df, cols, case, regions)
        opt.columns = base.columns
        d = (opt - base).loc[regs]
        return [(d[[k for k in keys if k in have]].sum(axis=1, min_count=1), color)
                for _, keys, color in COST_GROUPS]

    n = len(regs)
    xs = np.arange(n)
    fig, (axA, axB) = plt.subplots(
        2, 1, figsize=(13.5, 8.4), sharex=True,
        gridspec_kw={"height_ratios": [1.0, 1.35], "hspace": 0.16})
    nets = {}
    for ax, case, letter, label, clip in ((axA, case_a, "a", LABEL_GA, False),
                                          (axB, case_b, "b", LABEL_SCRATCH, True)):
        parts = _grouped(case)
        add_region_bands(ax, n)
        ax.grid(axis="y", color=GRID, lw=0.7, zorder=0.5)
        ax.axhline(0, color=MUTED, lw=0.9, zorder=2)
        top, bot = _stacked_change(ax, xs, parts)
        net = sum(np.asarray(v, float) for v, _ in parts)
        nets[case] = pd.Series(net, index=regs)
        ax.scatter(xs, net, marker="D", s=30, color=INK, edgecolors="white",
                   linewidths=0.8, zorder=5)
        lo = bot.min() * 1.08
        cap = _outlier_cap(top) if clip else None
        ax.set_ylim(lo, cap if cap else max(top.max() * 1.08, 0.05 * abs(lo)))
        if cap:   # off-scale columns: clipped bar + up-arrow + the net value
            for x in xs[top > cap]:
                ax.annotate("", xy=(x, cap), xytext=(x, cap * 0.86),
                            arrowprops=dict(arrowstyle="-|>", color=INK, lw=1.2),
                            zorder=6, annotation_clip=False)
                ax.text(x + 0.36, cap * 0.97, f"net {net[x]:+.1f}", rotation=90,
                        ha="left", va="top", fontsize=9, color=INK, zorder=6)
        minor_ticks(ax)
        _panel_tag(ax, letter, label)
    for ax in (axA, axB):
        ax.set_ylabel(f"Cost change vs LOADMATCH\n({TAE_IT})\n(2023 US cents/kWh)")
    style_region_axis(axB, [region_label(r) for r in regs])
    axA.tick_params(axis="x", length=0)

    handles = [plt.Rectangle((0, 0), 1, 1, color=c, label=g) for g, _, c in COST_GROUPS]
    handles.append(plt.Line2D([], [], marker="D", ls="", color=INK, mec="white",
                              ms=6.5, label="Net change"))
    axA.legend(handles=handles, loc="lower right", bbox_to_anchor=(1.0, 1.08),
               ncols=3, handlelength=1.2, columnspacing=1.2, borderaxespad=0.0)
    for ext in ("pdf", "png"):
        fig.savefig(outdir / f"fig_cost_change_by_category.{ext}")
    plt.close(fig)

    print("  wrote fig_cost_change_by_category.pdf/.png")
    print("    caption stats: cost groups (c/kWh, additive, ALL-NONH2-STORAGE excluded): "
          + "; ".join(f"{g} = {' + '.join(k)}" for g, k, _ in COST_GROUPS))
    for case, label in ((case_a, LABEL_GA), (case_b, LABEL_SCRATCH)):
        s_ = nets[case]
        print(f"    {plain_label(label)}: net change {s_.min():+.2f} to {s_.max():+.2f} "
              f"c/kWh ({int((s_ < 0).sum())} of {len(s_)} regions cheaper).")
    capped = [r for r, v in zip(regs, np.asarray(nets[case_b])) if axB.get_ylim()[1] < v]
    if capped:
        print(f"    panel b y-axis capped; off-scale: {capped}")


def fig_generation_mix_change(df, outdir, regions, case_a, case_b):
    # Share denominator = ALL generation sources (shares of total output).
    cols_all = _cols_in_group(df, G_GEN, exclude=GEN_EXCLUDE)

    def share(case):
        m = _case_matrix(df, cols_all, case, regions)
        m.columns = [_short(c) for c in cols_all]
        return m.div(m.sum(axis=1), axis=0) * 100.0

    sh = {c: share(c) for c in (BASELINE_CASE, case_a, case_b)}
    srcs = [s_ for s_ in GEN_SHOWN if s_ in sh[BASELINE_CASE].columns]
    if not srcs:
        print("  [SKIP] fig_generation_mix_change: no wind/solar generation columns.")
        return
    regs = [r for r in regions if np.isfinite(sh[BASELINE_CASE].loc[r, srcs]).all()]
    n = len(regs)
    xs = list(range(n))

    fig, axes = plt.subplots(len(srcs), 1, figsize=(13.5, 3.5 * len(srcs) + 1.3),
                             sharex=True, gridspec_kw={"hspace": 0.12})
    axes = np.atleast_1d(axes)
    for i, (ax, src) in enumerate(zip(axes, srcs)):
        add_region_bands(ax, n)
        ax.grid(axis="y", color=GRID, lw=0.7, zorder=0.5)
        h = _dumbbell(ax, xs, [
            (sh[BASELINE_CASE].loc[regs, src], C_BASELINE, "o", LABEL_BASELINE, SZ_BASE),
            (sh[case_a].loc[regs, src], C_GA, "o", LABEL_GA, SZ_GA),
            (sh[case_b].loc[regs, src], C_SCRATCH, "o", LABEL_SCRATCH, SZ_SCRATCH),
        ])
        ax.set_ylim(0, 100)
        ax.set_ylabel(f"{src} share of total\ngeneration (%)")
        minor_ticks(ax)
        ax.text(0.0, 1.02, "ab"[i], transform=ax.transAxes, fontweight="bold",
                fontsize=14, va="bottom")
    order_series = [LABEL_BASELINE, LABEL_GA, LABEL_SCRATCH]
    axes[0].legend([h[k] for k in order_series], order_series, loc="lower right",
                   bbox_to_anchor=(1.0, 1.0), ncols=3, handletextpad=0.2,
                   columnspacing=1.1, borderaxespad=0.2)
    style_region_axis(axes[-1], [region_label(r) for r in regs])
    for ax in axes[:-1]:
        ax.tick_params(axis="x", length=0)
    for ext in ("pdf", "png"):
        fig.savefig(outdir / f"fig_generation_mix_change.{ext}")
    plt.close(fig)

    print("  wrote fig_generation_mix_change.pdf/.png")
    for src in srcs + ["Solar heat"]:
        if src not in sh[BASELINE_CASE].columns:
            continue
        parts = []
        for case, label in ((case_a, LABEL_GA), (case_b, LABEL_SCRATCH)):
            dpp = (sh[case].loc[regs, src] - sh[BASELINE_CASE].loc[regs, src])
            parts.append(f"{plain_label(label)} {dpp.min():+.1f} to {dpp.max():+.1f} pp "
                         f"(median |change| {dpp.abs().median():.1f} pp)")
        print(f"    caption stats: {src} share change vs trial-and-error - "
              + "; ".join(parts) + ".")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--outdir", type=Path, default=RESULTS_DIR)
    ap.add_argument("--case-a", default=DEFAULT_CASE_A,
                    help="GA-from-trial-and-error case, blue (default %(default)r).")
    ap.add_argument("--case-b", default=DEFAULT_CASE_B,
                    help="GA-from-scratch case, orange (default %(default)r).")
    args = ap.parse_args(argv)
    args.outdir.mkdir(parents=True, exist_ok=True)

    cmp_df = pd.read_csv(args.csv)
    df = load_structure(args.xlsx)
    have_r = set(df["Identification :: Region"])
    have_c = set(df["Identification :: Case"])
    regions = [r for r in region_order(cmp_df) if r in have_r]
    if not regions:
        print("No regions found in both the CSV and the XLSX.")
        return
    for case in (args.case_a, args.case_b):
        if case not in have_c:
            print(f"  [WARN] case {case!r} not in the XLSX (available: {sorted(have_c)}).")
    # flag regions with no feasible baseline (change is undefined -> dropped)
    bl = df[df["Identification :: Case"] == BASELINE_CASE].set_index("Identification :: Region")
    cost = pd.to_numeric(bl["Annual cost ($B/yr) :: Cost MN ($B/yr)"], errors="coerce")
    infeasible = [r for r in regions if r in cost.index and not np.isfinite(cost.get(r))]
    if infeasible:
        print(f"  [INFO] baseline infeasible (shown grey / dropped): {infeasible}")

    fig_factor_changes(df, args.outdir, regions, args.case_a, args.case_b)
    fig_cost_change_by_category(df, args.outdir, regions, args.case_a, args.case_b)
    fig_generation_mix_change(df, args.outdir, regions, args.case_a, args.case_b)


if __name__ == "__main__":
    main()
